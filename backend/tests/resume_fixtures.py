"""Locally generated resume documents for the upload/extraction tests.

Nothing here downloads or embeds a real resume or calls an external
extraction API: PDFs are assembled byte by byte and DOCX files are written
with ``python-docx``, so the fixtures are deterministic, tiny and safe to
commit. The PDF writer below is the minimal subset needed for ``pypdf`` to
extract the text again.
"""

from __future__ import annotations

from io import BytesIO

from docx import Document as DocxDocument
from pypdf import PdfReader, PdfWriter


def _escape_pdf_text(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf_bytes(*lines: str) -> bytes:
    """Build a minimal single-page PDF whose text ``pypdf`` can extract.

    Each positional argument becomes one text-showing line. Called with no
    arguments it produces a page with no text at all, which stands in for a
    scanned/image-only PDF.
    """
    content_ops = ["BT", "/F1 14 Tf", "72 720 Td"]
    for index, line in enumerate(lines):
        if index:
            content_ops.append("0 -20 Td")
        content_ops.append(f"({_escape_pdf_text(line)}) Tj")
    content_ops.append("ET")
    stream = "\n".join(content_ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    buffer = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(buffer))
        buffer += f"{number} 0 obj\n".encode("ascii") + body + b"\nendobj\n"

    xref_offset = len(buffer)
    buffer += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    buffer += b"0000000000 65535 f \n"
    for offset in offsets:
        buffer += f"{offset:010d} 00000 n \n".encode("ascii")
    buffer += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode("ascii")
    return bytes(buffer)


def make_encrypted_pdf_bytes(*lines: str) -> bytes:
    """Build a password-protected PDF from the same minimal page builder."""
    reader = PdfReader(BytesIO(make_pdf_bytes(*lines)))
    writer = PdfWriter()
    writer.append(reader)
    writer.encrypt("secret-password")
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def make_docx_bytes(
    paragraphs: list[str], table: list[list[str]] | None = None
) -> bytes:
    """Build a DOCX with the given paragraphs and optional table rows."""
    document = DocxDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        rows = len(table)
        columns = max(len(row) for row in table)
        docx_table = document.add_table(rows=rows, cols=columns)
        for row_index, row in enumerate(table):
            for column_index, value in enumerate(row):
                docx_table.cell(row_index, column_index).text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = [
    "make_docx_bytes",
    "make_encrypted_pdf_bytes",
    "make_pdf_bytes",
]
